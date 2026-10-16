"""Generador de robots autónomos con IA (Gemini) basado en los estándares de .agents/"""

import os
import re
import requests
from typing import Dict, List, Optional, Tuple, Any

from core.config import DEFAULT_GEMINI_API_KEY


def inspect_sample_file(file_path: str, page_number: int = 1, target_sheet: str = "") -> str:
    """Inspecciona quirúrgicamente un archivo de muestra (Excel, PDF, CSV) para extraer su estructura y encabezados."""
    if not os.path.isfile(file_path):
        return f"El archivo '{file_path}' no existe en disco."

    ext = os.path.splitext(file_path)[1].lower()
    summary = []
    summary.append(f"Archivo: {os.path.basename(file_path)} (extensión: {ext})")

    try:
        if ext in (".xlsx", ".xls", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
            sheet_names = wb.sheetnames
            summary.append(f"Hojas disponibles en el archivo: {sheet_names}")

            # Seleccionar la hoja correcta
            selected_sheet_name = None
            if target_sheet:
                # Búsqueda exacta o parcial
                for s in sheet_names:
                    if target_sheet.lower() in s.lower() or s.lower() in target_sheet.lower():
                        selected_sheet_name = s
                        break

            if not selected_sheet_name:
                # Si la primera hoja es una portada (INICIO, PORTADA, INDICE), buscar la primera con datos reales
                first_lower = sheet_names[0].lower()
                if first_lower in ("inicio", "portada", "indice", "índice") and len(sheet_names) > 1:
                    idx = page_number if 0 < page_number < len(sheet_names) else 1
                    selected_sheet_name = sheet_names[idx]
                else:
                    idx = page_number - 1 if 0 <= (page_number - 1) < len(sheet_names) else 0
                    selected_sheet_name = sheet_names[idx]

            sheet = wb[selected_sheet_name]
            summary.append(f"Hoja analizada: '{selected_sheet_name}' (Total filas: {sheet.max_row}, Total cols: {sheet.max_column})")

            rows_data = []
            has_horizontal_months = False
            is_percentage_variation = False
            months_keywords = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

            for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if r_idx > 25:
                    break
                row_vals = [str(c).strip() if c is not None else "" for c in row]
                while row_vals and row_vals[-1] == "":
                    row_vals.pop()
                if any(row_vals):
                    row_str = " | ".join(row_vals[:14])
                    rows_data.append(f"Fila {r_idx:2d}: {row_str}")

                    # Detectar si hay meses en las columnas de encabezado
                    row_lower = row_str.lower()
                    month_matches = sum(1 for m in months_keywords if m in row_lower)
                    if month_matches >= 3:
                        has_horizontal_months = True
                    if "variación porcentual" in row_lower or "porcentaje" in row_lower or "índice" in row_lower or "indice" in row_lower:
                        is_percentage_variation = True

            wb.close()
            summary.append("\n".join(rows_data))

            if has_horizontal_months:
                summary.append(
                    "\n🚨 ATENCIÓN ARQUITECTURA: Se detectó una tabla HORIZONTAL/TRANSPUESTA. "
                    "Los Años y Meses están distribuidos en las columnas de encabezado, y los conceptos (ciudades o cuentas) están en las filas. "
                    "El robot DEBE iterar las columnas de fechas y generar registros [nv1, fecha, valor] des-pivotados."
                )

            summary.append(
                "📋 EXTRACCIÓN DE TÍTULOS DE ENCABEZADO (metadata_dict['titles']): "
                "Observa las filas superiores (filas 1 a 5 antes de la tabla de datos). "
                "Todas las líneas descriptivas presentes en el encabezado (ej: Fila 2: 'Cuadro Nº 1.2', "
                "Fila 3: 'BOLIVIA: VARIACIÓN PORCENTUAL...', Fila 4: '(Porcentaje)') DEBEN capturarse en orden como "
                "elementos separados de la lista 'titles' (ej: titles = [ws.cell(r, 1).value for r in range(2, 5) if ws.cell(r, 1).value]), "
                "de modo que insert_metadata cree limpiamente las columnas 'titulo1', 'titulo2', 'titulo3' en la tabla final."
            )

            if is_percentage_variation:
                summary.append(
                    "📌 NOTA DE VALIDACIÓN: El reporte contiene variaciones porcentuales o índices ponderados. "
                    "En validate_data_results(), NO realices suma aditiva de los componentes hacia el Total General (porque los porcentajes ponderados no se suman). "
                    "Valida que no esté vacío, columnas requeridas, tipo float y sin NaNs."
                )

        elif ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                total_pages = len(pdf.pages)
                summary.append(f"Total páginas PDF: {total_pages}")
                page_idx = page_number - 1 if 0 <= (page_number - 1) < total_pages else 0
                page = pdf.pages[page_idx]
                text = page.extract_text() or ""
                summary.append(f"Texto extraído de la página {page_idx + 1} (primeros 2000 caracteres):\n")
                summary.append(text[:2000])

        elif ext == ".csv":
            import pandas as pd
            df = pd.read_csv(file_path, nrows=15)
            summary.append(f"Columnas: {list(df.columns)}")
            summary.append(df.to_string(index=False))

        else:
            summary.append(f"Formato '{ext}' no reconocido para pre-inspección.")

    except Exception as exc:
        summary.append(f"Error durante la inspección del archivo: {exc}")

    return "\n".join(summary)


def load_datax_rules(workspace_dir: str = "") -> str:
    """Carga el texto de las reglas oficiales de .agents/rules/datax_standards.md."""
    if not workspace_dir:
        workspace_dir = os.path.dirname(os.path.dirname(__file__))
    rules_path = os.path.join(workspace_dir, ".agents", "rules", "datax_standards.md")
    if os.path.isfile(rules_path):
        try:
            with open(rules_path, "r", encoding="utf-8") as f:
                return f.read()
        except Exception:
            pass
    return ""


def load_doc(doc_name: str, workspace_dir: str = "") -> str:
    """Carga una guía de la carpeta docs/."""
    if not workspace_dir:
        workspace_dir = os.path.dirname(os.path.dirname(__file__))
    path = os.path.join(workspace_dir, "docs", doc_name)
    if os.path.isfile(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        except Exception:
            pass
    return ""


def load_example(example_subfolder: str, workspace_dir: str = "") -> str:
    """Carga el código de un robot de ejemplo oficial desde examples/<subfolder>/robot.py."""
    if not workspace_dir:
        workspace_dir = os.path.dirname(os.path.dirname(__file__))
    path = os.path.join(workspace_dir, "examples", example_subfolder, "robot.py")
    if os.path.isfile(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        except Exception:
            pass
    return ""


class AgentGenerator:
    """Motor de IA para generación y corrección de robots DataX."""

    def __init__(self, api_key: str = DEFAULT_GEMINI_API_KEY):
        self.api_key = api_key
        # Modelos en orden de preferencia y alta disponibilidad (priorizando endpoints lite con cuota activa)
        self.models = [
            "gemini-flash-lite-latest",
            "gemini-3.1-flash-lite",
            "gemini-3.5-flash-lite",
            "gemini-3.1-flash-lite-preview",
            "gemini-3.5-flash",
            "gemini-3.7-flash",
            "gemini-3-flash-preview",
            "gemini-flash-latest",
        ]

    def _call_gemini(self, prompt: str, system_instruction: str = "") -> Tuple[bool, str]:
        """Envía un prompt a Gemini con reintentos y tolerancia a 503/429 sobre modelos disponibles."""
        import time

        if not self.api_key:
            return False, "GEMINI_API_KEY no está configurada en .env"

        last_error = ""
        for model in self.models:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={self.api_key}"
            payload: Dict[str, Any] = {
                "contents": [{"parts": [{"text": prompt}]}],
                "generationConfig": {
                    "temperature": 0.2,
                    "maxOutputTokens": 8192,
                }
            }
            if system_instruction:
                payload["systemInstruction"] = {
                    "parts": [{"text": system_instruction}]
                }

            try:
                res = requests.post(url, json=payload, timeout=45)
                if res.status_code == 200:
                    data = res.json()
                    candidates = data.get("candidates", [])
                    if candidates:
                        text_resp = candidates[0].get("content", {}).get("parts", [{}])[0].get("text", "")
                        return True, text_resp
                    return False, "Gemini no retornó candidatos de respuesta."
                elif res.status_code in (404, 429, 500, 503):
                    # Error temporal o cuota excedida en este modelo específico, probar siguiente modelo
                    last_error = f"{model}: {res.status_code} - {res.text[:120]}"
                    time.sleep(0.5)
                    continue
                else:
                    return False, f"Error {res.status_code} de Gemini: {res.text[:200]}"
            except Exception as exc:
                last_error = f"{model} excepción: {exc}"
                time.sleep(0.5)
                continue

        return False, f"No se pudo completar la llamada con ningún modelo. Último error: {last_error}"

    @staticmethod
    def _clean_code_block(text: str) -> str:
        """Extrae el bloque de código python limpio."""
        match = re.search(r"```python\s*(.*?)\s*```", text, re.DOTALL)
        if match:
            return match.group(1).strip()
        match_generic = re.search(r"```\s*(.*?)\s*```", text, re.DOTALL)
        if match_generic:
            return match_generic.group(1).strip()
        return text.strip()

    def generate_conversion_robot(
        self,
        robot_code: str,
        report_name: str,
        page_number: int = 1,
        decimal_separator: str = ",",
        key_words: str = "",
        sample_inspection: str = "",
        target_sheet: str = "",
        rules_text: str = ""
    ) -> Tuple[bool, str]:
        """Genera el código de un robot de conversión cumpliendo la Regla Estricta de 2 Métodos y estándares oficiales."""
        system_prompt = (
            "Eres el Ingeniero de Datos Senior y Tutor de DATAX Platform. "
            "Tu tarea es generar el archivo robot.py para un robot de CONVERSIÓN de datos financieros. "
            "Debes cumplir ESTRICTAMENTE con los estándares oficiales de DATAX:\n"
            "1. IMPORTS Y BASE CLASS OBLIGATORIA (ESTÁNDAR ESTRICTO DE DATAX PLATFORM):\n"
            "   from models.conversion.Conversion_Base import Conversion_Base\n"
            "   from models.conversion.tools.conversion_tools import to_numeric_datax\n"
            "   # REGLA ESTRICTA: NO uses bloques try/except en las importaciones. NO importes desde conversion_tools suelto. Usa directamente models.conversion.\n\n"
            "2. CLASE Y REGLA ESTRICTA DE 2 MÉTODOS:\n"
            f"   class {robot_code}(Conversion_Base):\n"
            "       def extraction(self, file_path: str, key_words: str, template_path: str = '', page_number: int = 1, format: str = '%Y-%m-%d') -> tuple:\n"
            "       def validate_data_results(self, dataframe: pd.DataFrame, decimal_separator: str = ',', TOLERANCE: float = 6.0) -> bool:\n\n"
            "3. PROHIBIDO crear __init__ o métodos auxiliares en self. Funciones auxiliares específicas van anidadas dentro de extraction().\n"
            "4. PROHIBIDO hardcodear listas de nombres de cuentas o rubros contables. Usa celdas pivote, indentaciones, sangrías o signos aritméticos.\n"
            "5. extraction() retorna (metadata_dict, df_melted) donde:\n"
            "   - metadata_dict DEBE contener única y exactamente estas 3 llaves obligatorias:\n"
            "       'file_name': os.path.basename(file_path),\n"
            "       'titles': [sheet_title],  # DEBE SER UNA LISTA DE STRINGS (ej: ['Titulo 1', 'Titulo 2'])\n"
            "       'page_number': int(page_number)\n"
            "   - df_melted contiene las columnas limpias des-pivotadas: [nv1, nv2, ..., fecha, valor].\n"
            "   - NO llamar a insert_metadata() dentro de extraction() porque el orquestador y test_conversion.py lo invocan sobre el resultado.\n"
            "6. validate_data_results() retorna False cuando los datos son válidos y consistentes, y True si hay error.\n"
            f"7. Al final del archivo incluir alias: Robot = {robot_code}.\n"
            "8. Responder ÚNICAMENTE con el bloque de código Python limpio dentro de ```python ... ```."
        )

        # Cargar guía y ejemplo de referencia oficial
        is_pdf = ".pdf" in (sample_inspection or "").lower() or (key_words or "").lower().endswith(".pdf")
        example_subfolder = "conversion_pdf" if is_pdf else "conversion_excel"
        example_code = load_example(example_subfolder)
        conversion_guide = load_doc("GUIA_ROBOTS_CONVERSION.md")

        user_prompt = f"""
Genera el código completo de robot.py para el siguiente reporte de Conversión:
- Código de reporte: {robot_code}
- Nombre del reporte: {report_name}
- Página sugerida: {page_number}
- Hoja objetivo: {target_sheet or 'Analizar la muestra adjunta para identificar la hoja correcta'}
- Separador decimal: {decimal_separator}
- Palabras clave: {key_words}

Estructura / Muestra quirúrgica del archivo:
{sample_inspection or 'No se proporcionó muestra previa. Genera la estructura basada en los estándares con openpyxl/pandas.'}

INSTRUCCIONES CLAVE DE ARQUITECTURA Y PRUEBAS (test_conversion.py):
1. Importaciones oficiales exactas (sin try/except):
   from models.conversion.Conversion_Base import Conversion_Base
   from models.conversion.tools.conversion_tools import to_numeric_datax
2. metadata_dict DEBE contener exacta y obligatoriamente estas 3 llaves:
   {{
       'file_name': os.path.basename(file_path),
       'titles': [str(title)],
       'page_number': int(page_number)
   }}
3. Revisa atentamente si la muestra analizada indica que la tabla es HORIZONTAL (Años y Meses en columnas de encabezado, conceptos como ciudades o cuentas en filas) o VERTICAL (fechas en filas).
   - Si es HORIZONTAL: el robot debe iterar sobre las columnas de fechas (construyendo fechas YYYY-MM-DD del último día de cada mes con datetime o last_day_month) y generar registros [nv1, fecha, valor] des-pivotados.
4. Si el reporte es de Variación Porcentual, Índices o Tasas (no aditivos), validate_data_results() NO debe sumar los componentes individuales contra el total (porque los porcentajes ponderados no se suman); valida que el DataFrame no esté vacío, contenga las columnas obligatorias, valores numéricos válidos y sin NaNs.
5. Devuelve False en validate_data_results si los datos son válidos y consistentes (test_conversion.py fallará si retorna True).
6. NO llames a insert_metadata() dentro de extraction(). Deja el DataFrame con [nv1, fecha, valor] y retorna metadata_dict, df_melted.

EJEMPLO OFICIAL DE REFERENCIA ({example_subfolder}):
```python
{example_code[:2500] if example_code else '# Usar arquitectura oficial de 2 métodos'}
```

GUÍA OFICIAL DE CONVERSIÓN:
{conversion_guide[:1500] if conversion_guide else ''}

Información adicional de reglas y estándares:
{rules_text[:2000] if rules_text else 'Cumplir estrictamente con PEP 8 y estándares oficiales de DATAX.'}
"""
        ok, res = self._call_gemini(user_prompt, system_instruction=system_prompt)
        if ok:
            return True, self._clean_code_block(res)
        return False, res

    def generate_download_robot(
        self,
        robot_code: str,
        download_type: str,
        main_url: str,
        updated_to: str = "2024-01-01",
        report_name: str = "",
        rules_text: str = ""
    ) -> Tuple[bool, str]:
        """Genera el código de un robot de descarga según su tipo (Tipo I, II, III o IV)."""
        system_prompt = (
            "Eres el Ingeniero de Datos Senior de DATAX Platform especializado en robots de DESCARGA (Web Scraping y APIs). "
            "Genera el archivo robot.py cumpliendo estrictamente con los estándares oficiales de DATAX:\n"
            "1. IMPORTS Y BASE CLASS OBLIGATORIA (ESTÁNDAR ESTRICTO DE DATAX PLATFORM):\n"
            "   from models.download.Download_Base import Download_Base\n"
            "   from models.download.tools.download_tools import (\n"
            "       format_date, month_to_number, month_abr_to_number, last_day_month, createDirectoryStruct, remove_tmp, text_extract\n"
            "   )\n"
            "   # REGLA ESTRICTA: NO uses bloques try/except en las importaciones. Usa directamente models.download.\n\n"
            f"2. Nombre de clase igual al código del robot: class {robot_code}(Download_Base):\n"
            f"   Al final del archivo alias: Robot = {robot_code}.\n"
            "3. Estructura exacta según el Tipo de descarga solicitado:\n"
            "   - Tipo I: métodos get_file_url(self, main_url, updated_to, key_words='', format='%Y-%m-%d') -> list y compare_files(self, files_paths, updated_to, format='%Y-%m-%d') -> list\n"
            "   - Tipo II: métodos verify_download(self, main_url, updated_to, path, key_words='', format='%Y-%m-%d') -> list y compare_files(self, files_paths, updated_to, format='%Y-%m-%d') -> list\n"
            "   - Tipo III: método check_new_data(self, main_url, updated_to, path, key_words=None, format='%Y-%m-%d') -> list\n"
            "   - Tipo IV: método get_data(self, main_url, path, updated_to, format='%Y-%m-%d') -> str\n"
            "4. Manejar siempre Playwright en modo headless si se requiere navegador con bloques try...finally asegurando browser.close().\n"
            "5. Limpiar siempre la carpeta tmp con shutil.rmtree y os.makedirs(tmp, exist_ok=True).\n"
            "6. Filtrar estrictamente fechas mayores a updated_to (file_date > updated_to).\n"
            "7. Llaves exactas en los diccionarios de retorno: 'tmp_path', 'updated_to', 'download_url'.\n"
            "8. Responder ÚNICAMENTE con el bloque de código Python limpio dentro de ```python ... ```."
        )

        type_map = {
            "Tipo I": "download_type_i",
            "Tipo II": "download_type_ii",
            "Tipo III": "download_type_iii",
            "Tipo IV": "download_type_iv",
        }
        example_subfolder = type_map.get(download_type, "download_type_i")
        example_code = load_example(example_subfolder)
        playwright_example = load_example("download_playwright")
        download_guide = load_doc("GUIA_ROBOTS_DESCARGA.md")

        user_prompt = f"""
Genera el código completo de robot.py para el siguiente robot de Descarga:
- Código del robot: {robot_code}
- Nombre: {report_name}
- Tipo de Descarga: {download_type}
- URL principal: {main_url}
- Fecha de referencia (updated_to): {updated_to}

EJEMPLO OFICIAL DE REFERENCIA ({download_type}):
```python
{example_code if example_code else '# Usar contrato oficial de ' + download_type}
```

REFERENCIA DE PLAYWRIGHT HEADLESS:
```python
{playwright_example[:1200] if playwright_example else ''}
```

GUÍA OFICIAL DE DESCARGA:
{download_guide[:1500] if download_guide else ''}

Detalles de estándares y reglas:
{rules_text[:2000] if rules_text else 'Usar download_tools (createDirectoryStruct, month_to_number, text_extract).'}
"""
        ok, res = self._call_gemini(user_prompt, system_instruction=system_prompt)
        if ok:
            return True, self._clean_code_block(res)
        return False, res

    def heal_robot_code(
        self,
        current_code: str,
        error_message: str,
        robot_code: str,
        task_type: str = "conversion"
    ) -> Tuple[bool, str]:
        """Corrige el código de un robot tras una falla en la ejecución o validación."""
        system_prompt = (
            "Eres el Ingeniero de Datos Senior y Tutor de DATAX Platform. "
            "El código del robot falló en la prueba de validación unitaria. Tu objetivo es CORREGIR el código Python cumpliendo estrictamente los estándares oficiales:\n"
            "1. Las importaciones deben ser ESTRICTAMENTE las oficiales de DATAX Platform, sin try/except ni fallbacks:\n"
            "   Para Conversión:\n"
            "       from models.conversion.Conversion_Base import Conversion_Base\n"
            "       from models.conversion.tools.conversion_tools import to_numeric_datax\n"
            "   Para Descarga:\n"
            "       from models.download.Download_Base import Download_Base\n"
            "       from models.download.tools.download_tools import format_date, month_to_number, last_day_month\n"
            "2. Para Conversión: ÚNICA Y EXCLUSIVAMENTE dos métodos: extraction() y validate_data_results().\n"
            "3. En extraction(), metadata_dict DEBE contener exactamente las 3 llaves: 'file_name', 'titles' (como lista de strings) y 'page_number' (como int).\n"
            "4. NO llamar a insert_metadata() dentro de extraction().\n"
            "5. En validate_data_results(): Retornar False si los datos son válidos, y True solo si hay errores reales.\n"
            "6. Responde ÚNICAMENTE con el código Python corregido completo dentro de ```python ... ```."
        )

        checklist = load_doc("CHECKLIST_ENTREGA.md")

        user_prompt = f"""
El robot {robot_code} ({task_type}) produjo el siguiente error durante la ejecución:

ERROR / TRACEBACK:
{error_message}

CÓDIGO ACTUAL:
```python
{current_code}
```

CHECKLIST OFICIAL DE CALIDAD:
{checklist[:1500] if checklist else ''}

Analiza la causa raíz del error (ej. índice fuera de rango, tipos de datos, jerarquía rota, método no encontrado, etc.) y genera la versión corregida completa del archivo robot.py.
"""
        ok, res = self._call_gemini(user_prompt, system_instruction=system_prompt)
        if ok:
            return True, self._clean_code_block(res)
        return False, res
